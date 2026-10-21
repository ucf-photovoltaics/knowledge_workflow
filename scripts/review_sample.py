"""Blind expert review of a random sample of finished ontologies: four checks, one per reviewer.

    uv run --with-requirements requirements.txt --with openpyxl python scripts/review_sample.py sample --name models-n25 \
        --sweeps pv-cost-chain-llama31-8b pv-cost-chain-gptoss-120b [--runs RUN ...] [--per-run 20] [--seed 1]
    uv run --with-requirements requirements.txt --with openpyxl python scripts/review_sample.py score --name models-n25

sample writes outputs/review/<name>/:
  placement.xlsx    reviewer 1: is every <term> a <parent>?
  definitions.xlsx  reviewer 2: is the definition correct, and supported by the paper excerpts?
  relations.xlsx    reviewer 3: is the relation between two terms, or the mapping to an external ontology, correct?
  values.xlsx       reviewer 4: is the reported value (property, technology, value, unit, basis) what the quote says?
  key.csv           item -> run, model, domain and origin. Keep it from the reviewers.
Items are drawn at random from each run's final ontology (the same number per run, reproducible from --seed) and
shuffled across runs; workbooks show no run, model or IRI. score reads the returned workbooks and writes scores.csv:
per check and model (and per origin), the share judged correct with a 95% Wilson interval.
"""
import argparse
import csv
import json
import math
import random
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.config import OUTPUTS  # noqa: E402
from src.tools import ontostore  # noqa: E402

CHECKS = {
    "placement": {
        "question": "Is every <Term> a kind of <Parent>? (Is-a, not part-of or related-to.)",
        "answers": ["Correct", "Correct but too general", "Wrong", "Not a valid term", "Can't judge"],
        "correct": {"Correct"}, "acceptable": {"Correct", "Correct but too general"},
        "notes": ["Correct: every instance of the term is an instance of the parent, and no clearly better parent exists.",
                  "Correct but too general: true, but a more specific parent exists; name it under Better parent.",
                  "Wrong: some or all instances of the term are not instances of the parent.",
                  "Not a valid term: not a meaningful concept in this domain (fragment, typo, too vague).",
                  "Can't judge: outside your expertise. These items are left out of the scores."]},
    "definitions": {
        "question": "Is the Definition correct for the Term, and supported by the Paper excerpts?",
        "answers": ["Correct and supported", "Correct, not in excerpts", "Partly correct", "Wrong", "Not a valid term",
                    "Can't judge"],
        "correct": {"Correct and supported"}, "acceptable": {"Correct and supported", "Correct, not in excerpts"},
        "notes": ["Correct and supported: accurate, and the excerpts say the same thing.",
                  "Correct, not in excerpts: accurate, but the excerpts don't support it (added knowledge).",
                  "Partly correct: the main idea is right but some of it is wrong, too narrow or too broad.",
                  "Wrong: inaccurate, or it defines a different concept.",
                  "Not a valid term: not a meaningful concept in this domain.",
                  "Can't judge: outside your expertise. These items are left out of the scores."]},
    "relations": {
        "question": "Is the stated Relation between Term and Target correct? (Use the evidence quote where given.)",
        "answers": ["Correct", "Right terms, wrong relation", "Wrong", "Can't judge"],
        "correct": {"Correct"}, "acceptable": {"Correct"},
        "notes": ["Correct: the relation holds as stated, in that direction.",
                  "Right terms, wrong relation: the two are related, but not as stated (e.g. exact vs broader, causes "
                  "vs part of, wrong direction).",
                  "Wrong: the two are not related this way at all.",
                  "Can't judge: outside your expertise. These items are left out of the scores."]},
    "values": {
        "question": "Is the reported Value right: the Property, the Technology it is for, the value and unit, and the "
                    "Basis, as the Quote states them?",
        "answers": ["Correct", "Correct, technology missing", "Wrong technology", "Wrong property", "Wrong value or unit",
                    "Wrong basis", "Not stated in quote", "Can't judge"],
        "correct": {"Correct"}, "acceptable": {"Correct", "Correct, technology missing"},
        "notes": ["Correct: everything shown matches the quote (Technology may be blank when the quote names none).",
                  "Correct, technology missing: right, but the quote names the device or technology and it is blank.",
                  "Wrong technology: the value is for something else than the Technology shown.",
                  "Wrong property: the value is not a value of this Property.",
                  "Wrong value or unit: the number, range or unit differs from the quote.",
                  "Wrong basis: measured, simulated, cited or target is wrong (e.g. a cited value shown as measured).",
                  "Not stated in quote: the quote does not give this value.",
                  "Can't judge: outside your expertise. These items are left out of the scores."]},
}
BASIS_WORDS = {"measured": "measured in this work", "simulated": "simulated or calculated in this work",
               "cited": "cited from another source", "target": "target, projection or assumed input"}
COLUMNS = {
    "placement": ["Item", "Term", "Term definition", "Found in", "Parent", "Parent definition", "Parent from", "Answer",
                  "Better parent", "Notes"],
    "definitions": ["Item", "Term", "Definition", "Paper excerpts", "Found in", "Answer", "Notes"],
    "relations": ["Item", "Kind", "Term", "Term definition", "Found in", "Relation", "Target", "Target definition", "Evidence",
                  "Answer", "Notes"],
    "values": ["Item", "Property", "Technology", "Value", "Unit", "Condition", "Basis", "Quote", "Found in", "Answer",
               "Notes"],
}
MAPPING_WORDS = {"equivalent": "is equivalent to", "exact": "is the same as", "close": "is nearly the same as",
                 "broader": "is a kind of", "narrower": "is more general than", "related": "is related to"}


def short(text, n=400):
    text = " ".join(str(text or "").split())
    return text if len(text) <= n else text[:n - 1] + "…"


def load_runs(names):
    runs = []
    for name in names:
        d = OUTPUTS / name
        manifest = json.loads((d / "run.json").read_text(encoding="utf-8-sig"))
        classes = json.loads((d / "ontology" / "enriched.json").read_text(encoding="utf-8"))
        mappings = json.loads((d / "ontology" / "mappings.json").read_text(encoding="utf-8"))
        titles = {}
        for p in (d / "papers").glob("*.json"):
            titles[p.stem] = json.loads(p.read_text(encoding="utf-8")).get("title", p.stem)
        cfg = manifest.get("config", {})
        values = d / "normalized" / "measurements.json"
        labels = {c["id"]: c["label"] for c in json.loads((d / "normalized" / "concepts.json").read_text(encoding="utf-8"))}
        runs.append({"id": name, "domain": manifest.get("collection", {}).get("name", ""), "cfg": cfg,
                     "classes": [c for c in classes if not c.get("excluded")], "mappings": mappings, "titles": titles,
                     "labels": labels,
                     "values": json.loads(values.read_text(encoding="utf-8")) if values.exists() else []})
    return runs


def sweep_runs(sweep):
    """Completed domain runs of a sweep, at its largest size."""
    rows = list(csv.DictReader(open(OUTPUTS / f"sweep-{sweep}" / "summary.csv", encoding="utf-8")))
    done = [r for r in rows if r["kind"] == "domain" and r["status"] == "completed"]
    top = max((int(r["limit"]) for r in done), default=None)
    return [r["run_id"] for r in done if int(r["limit"]) == top]


def describe(iri, by_iri):
    """Label, definition and source of a parent or target: a class of this ontology or a store term."""
    if iri in by_iri:
        c = by_iri[iri]
        return c["label"], c.get("definition") or (c["definitions"][0]["text"] if c["definitions"] else ""), "this ontology"
    t = ontostore.term(iri) or {}
    return t.get("label") or iri.rsplit("/", 1)[-1], t.get("definition", ""), t.get("ontology", "")


def definition(c):
    return c.get("definition") or (c["definitions"][0]["text"] if c["definitions"] else "")


def found_in(run, c):
    """Titles of up to two papers the term comes from, for context."""
    return "\n".join(short(run["titles"].get(p, p), 120) for p in c.get("papers", [])[:2])


def placement(run, rng, n):
    by_iri = {c["iri"]: c for c in run["classes"]}
    items = []
    for c in rng.sample(run["classes"], min(n, len(run["classes"]))):
        label, text, source = describe(c["parent"], by_iri)
        items.append(({"Term": c["label"], "Term definition": short(definition(c)), "Found in": found_in(run, c),
                       "Parent": label,
                       "Parent definition": short(text), "Parent from": source},
                      {"model": run["cfg"].get("model_ontology", ""), "class_id": c["id"],
                       "origin": c.get("parent_source", "")}))
    return items


def definitions(run, rng, n):
    pool = [c for c in run["classes"] if c.get("definition")]
    items = []
    for c in rng.sample(pool, min(n, len(pool))):
        excerpts = "\n".join(f"[{short(run['titles'].get(d['paper'], d['paper']), 80)}] {short(d['text'], 300)}"
                             for d in c["definitions"][:3]) or "(no excerpt)"
        items.append(({"Term": c["label"], "Definition": short(c["definition"], 600), "Paper excerpts": excerpts,
                       "Found in": found_in(run, c)},
                      {"model": run["cfg"].get("model_enrichment", ""), "class_id": c["id"],
                       "origin": c.get("definition_status", "")}))
    return items


def relations(run, rng, n):
    by_id = {c["id"]: c for c in run["classes"]}
    by_iri = {c["iri"]: c for c in run["classes"]}
    seen, restrictions = set(), []
    for c in run["classes"]:  # one item per asserted link; a local property and its RO/CCO twin are one assertion
        for r in sorted(c.get("restrictions", []), key=lambda r: str(r.get("p", "")).startswith("local:")):
            key = (c["id"], r.get("o"), r.get("phrase"))
            if r.get("o") in by_id and key not in seen:
                seen.add(key)
                restrictions.append((c, r))
    mappings = [m for m in run["mappings"] if m.get("id") in by_id]
    half = n // 2
    items = []
    for c, r in rng.sample(restrictions, min(half, len(restrictions))):
        p = str(r.get("p", ""))
        relation = p.split(":", 1)[1] if p.startswith("local:") else describe(p, by_iri)[0]
        if r.get("phrase") and r["phrase"].replace("_", " ") not in relation:
            relation += f" ({r['phrase'].replace('_', ' ')})"
        ev = r.get("evidence") or []
        quote = (f"[{short(run['titles'].get(ev[0]['paper'], ev[0]['paper']), 80)}] {short(ev[0]['text'], 300)}"
                 if ev else "(no quote)")
        o = by_id[r["o"]]
        items.append(({"Kind": "Relation", "Term": c["label"], "Term definition": short(definition(c)),
                       "Found in": found_in(run, c),
                       "Relation": relation, "Target": o["label"], "Target definition": short(definition(o)),
                       "Evidence": quote},
                      {"model": run["cfg"].get("model_enrichment", ""), "class_id": c["id"],
                       "origin": f"relation:{r.get('kind', '')}"}))
    for m in rng.sample(mappings, min(n - half, len(mappings))):
        c = by_id[m["id"]]
        items.append(({"Kind": "Mapping", "Term": c["label"], "Term definition": short(definition(c)),
                       "Found in": found_in(run, c),
                       "Relation": MAPPING_WORDS.get(m.get("relation"), m.get("relation", "")),
                       "Target": f"{m.get('label', '')} ({m.get('ontology', '')})",
                       "Target definition": short(m.get("definition", "")), "Evidence": ""},
                      {"model": run["cfg"].get("model_interoperability", ""), "class_id": c["id"],
                       "origin": f"mapping:{m.get('relation', '')}"}))
    return items


def values(run, rng, n):
    label = run["labels"]
    items = []
    for m in rng.sample(run["values"], min(n, len(run["values"]))):
        items.append(({"Property": label.get(m.get("concept"), ""), "Technology": label.get(m.get("entity"), ""),
                       "Value": m.get("value", ""), "Unit": m.get("unit", ""), "Condition": short(m.get("condition"), 200),
                       "Basis": BASIS_WORDS.get(m.get("basis"), "(not given)"), "Quote": short(m.get("evidence"), 300),
                       "Found in": short(run["titles"].get(m.get("paper"), m.get("paper", "")), 120)},
                      {"model": run["cfg"].get("model_extraction", ""), "class_id": m.get("concept", ""),
                       "origin": f"basis:{m.get('basis') or 'none'}"}))
    return items


def write_book(path, check, rows):
    spec = CHECKS[check]
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    cols = COLUMNS[check]
    ws.append(cols)
    for row in rows:
        ws.append([row.get(c, "") for c in cols])
    widths = {"Item": 8, "Kind": 10, "Answer": 24, "Parent from": 14, "Relation": 24, "Found in": 30, "Value": 12,
              "Unit": 10, "Basis": 20}
    for i, c in enumerate(cols, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths.get(c, 40)
        ws.cell(1, i).font = Font(bold=True)
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    answer = ws.cell(1, cols.index("Answer") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(spec["answers"]) + '"', allow_blank=True)
    dv.add(f"{answer}2:{answer}{len(rows) + 1}")
    ws.add_data_validation(dv)
    info = wb.create_sheet("Instructions", 0)
    for line in [spec["question"], "", "Choose one Answer per row from the drop-down list:", *spec["notes"], "",
                 "Judge each row on its own, with your domain knowledge and the text shown. Rows come from several "
                 "ontologies in random order. Leave the Item column unchanged."]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    wb.save(path)


def sample(args):
    names = list(args.runs or [])
    for s in args.sweeps or []:
        names += sweep_runs(s)
    if not names:
        sys.exit("give --runs or --sweeps")
    ontostore.require()
    runs = load_runs(dict.fromkeys(names))
    out = OUTPUTS / "review" / args.name
    if out.exists() and any(out.glob("*.xlsx")) and not args.force:
        sys.exit(f"{out} already has workbooks; use --force to replace them (reviewers may be working on them)")
    out.mkdir(parents=True, exist_ok=True)
    key = []
    for check, draw in (("placement", placement), ("definitions", definitions), ("relations", relations),
                        ("values", values)):
        pool = []
        for run in runs:
            rng = random.Random(f"{args.seed}:{check}:{run['id']}")
            pool += [(row, {**k, "run_id": run["id"], "domain": run["domain"]}) for row, k in draw(run, rng, args.per_run)]
        random.Random(f"{args.seed}:{check}:order").shuffle(pool)
        prefix = check[0].upper()
        for i, (row, k) in enumerate(pool, 1):
            row["Item"] = f"{prefix}{i:04d}"
            key.append({"check": check, "item": row["Item"], **k})
        write_book(out / f"{check}.xlsx", check, [row for row, _ in pool])
        print(f"{check}: {len(pool)} items from {len(runs)} runs -> {out / (check + '.xlsx')}")
    with open(out / "key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, ["check", "item", "run_id", "model", "domain", "class_id", "origin"])
        w.writeheader()
        w.writerows(key)
    print(f"key: {out / 'key.csv'} (do not share with reviewers)")


def wilson(k, n, z=1.96):
    if not n:
        return 0.0, 0.0
    p = k / n
    centre, half = (p + z * z / (2 * n)) / (1 + z * z / n), z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return max(0.0, centre - half), min(1.0, centre + half)


def score(args):
    out = OUTPUTS / "review" / args.name
    key = {(r["check"], r["item"]): r for r in csv.DictReader(open(out / "key.csv", encoding="utf-8"))}
    groups = defaultdict(lambda: defaultdict(int))
    unanswered = defaultdict(int)
    for check, spec in CHECKS.items():
        path = out / f"{check}.xlsx"
        if not path.exists():
            continue
        ws = load_workbook(path, read_only=True)["Review"]
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = dict(zip(header, values))
            k = key.get((check, row.get("Item")))
            answer = (row.get("Answer") or "").strip()
            if not k:
                continue
            if not answer:
                unanswered[check] += 1
                continue
            if answer == "Can't judge":
                continue
            for group in ((check, k["model"], "all"), (check, k["model"], k["origin"])):
                g = groups[group]
                g["n"] += 1
                g["correct"] += answer in spec["correct"]
                g["acceptable"] += answer in spec["acceptable"]
                g["invalid_term"] += answer == "Not a valid term"
    rows = []
    for (check, model, origin), g in sorted(groups.items()):
        lo, hi = wilson(g["correct"], g["n"])
        alo, ahi = wilson(g["acceptable"], g["n"])
        rows.append({"check": check, "model": model, "origin": origin, "judged": g["n"],
                     "correct": round(g["correct"] / g["n"], 3), "correct_ci95": f"{lo:.2f}-{hi:.2f}",
                     "acceptable": round(g["acceptable"] / g["n"], 3), "acceptable_ci95": f"{alo:.2f}-{ahi:.2f}",
                     "invalid_term": round(g["invalid_term"] / g["n"], 3)})
    with open(out / "scores.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, list(rows[0]) if rows else ["check"])
        w.writeheader()
        w.writerows(rows)
    for r in rows:
        if r["origin"] == "all":
            print(f"{r['check']:12} {r['model']:40} n={r['judged']:4}  correct {r['correct']:.0%} ({r['correct_ci95']})"
                  f"  acceptable {r['acceptable']:.0%}  invalid terms {r['invalid_term']:.0%}")
    for check, n in unanswered.items():
        print(f"{check}: {n} items not answered yet")
    print(f"scores: {out / 'scores.csv'} (per model, and per origin: how a parent, definition or link was produced)")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample")
    s.add_argument("--name", required=True, help="review folder under outputs/review/")
    s.add_argument("--runs", nargs="+", help="domain run ids")
    s.add_argument("--sweeps", nargs="+", help="sweep ids: their completed domain runs at the largest size")
    s.add_argument("--per-run", type=int, default=20, help="items per run and check (default 20)")
    s.add_argument("--seed", default="1")
    s.add_argument("--force", action="store_true")
    c = sub.add_parser("score")
    c.add_argument("--name", required=True)
    args = ap.parse_args()
    sample(args) if args.cmd == "sample" else score(args)


if __name__ == "__main__":
    main()
